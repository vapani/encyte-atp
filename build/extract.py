#!/usr/bin/env python3
"""Pull likely engagement fields out of a proposal (.docx or .pdf).

This is pattern matching, not comprehension. It finds things that have a shape -
an ABN is eleven digits, a reference looks like 26-NGA-WD-062, money starts with
a dollar sign - and it guesses at the rest. It is wrong often enough that every
field it returns must be reviewed before a contract is built from it.

Each field comes back as {"value": ..., "confidence": "high"|"low"} so the form
can show you which ones to look at hardest.
"""
import json, os, re, sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

MONTHS = ("January February March April May June July August September "
          "October November December").split()


def read_text(path):
    """Plain text of a .docx or .pdf. Raises with a readable message otherwise."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        from docx import Document
        d = Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext == ".pdf":
        import fitz
        with fitz.open(path) as doc:
            return "\n".join(p.get_text() for p in doc)
    raise ValueError(f"unsupported file type '{ext}' - upload a .docx or .pdf")


def _hi(v):  return {"value": v, "confidence": "high"}
def _lo(v):  return {"value": v, "confidence": "low"}


def extract(path):
    text = read_text(path)
    flat = re.sub(r"[ \t]+", " ", text)
    out, notes = {}, []

    provider = json.load(open(f"{ROOT}/jurisdictions/AU.json"))["provider"]
    own_abn = re.sub(r"\s", "", provider["abn"])
    own_name = provider["legal_name"].lower()

    # --- reference, e.g. 26-NGA-WD-062
    refs = re.findall(r"\b(\d{2}-[A-Z0-9]{2,8}-[A-Z]{2}-\d{2,4})\b", flat)
    if refs:
        out["proposal.ref"] = _hi(refs[0])
        out["atp.ref"] = _lo(refs[0])          # usually the same series, often not identical
        if len(set(refs)) > 1:
            notes.append(f"several references found ({', '.join(sorted(set(refs)))}) - first used")

    # --- dates written as 8 September 2026
    dates = re.findall(r"\b(\d{1,2} (?:%s) \d{4})\b" % "|".join(MONTHS), flat)
    if dates:
        out["proposal.date"] = (_hi if len(set(dates)) == 1 else _lo)(dates[0])
        if len(set(dates)) > 1:
            notes.append(f"{len(set(dates))} dates found - earliest-appearing used for the proposal date")

    # --- ABNs, ignoring your own
    abns = re.findall(r"\b(\d{2}[ ]?\d{3}[ ]?\d{3}[ ]?\d{3})\b", flat)
    others = [a for a in abns if re.sub(r"\s", "", a) != own_abn]
    if others:
        out["client.abn"] = (_hi if len(set(others)) == 1 else _lo)(others[0])

    # --- client legal name: a company/trust that is not you
    ents = re.findall(r"\b([A-Z][A-Za-z0-9&.,'\- ]{2,60}?(?:Pty Ltd|Pty\. Ltd\.|Limited|Ltd|Trust))\b", flat)
    cands = []
    for e in ents:
        e = e.strip(" ,.")
        if own_name.split()[0] in e.lower():
            continue
        if e.lower() not in [c.lower() for c in cands]:
            cands.append(e)
    if cands:
        out["client.legal_name"] = (_hi if len(cands) == 1 else _lo)(cands[0])
        short = re.split(r"\s+(?:Pty|Limited|Ltd)\b", cands[0])[0].strip()
        out["client.short_name"] = _lo(short)
        if len(cands) > 1:
            notes.append(f"possible client names: {', '.join(cands[:4])}")


    # --- money: the largest figure is usually the project price
    amounts = []
    for m in re.findall(r"\$\s?([\d,]+(?:\.\d{2})?)", flat):
        try:
            amounts.append(float(m.replace(",", "")))
        except ValueError:
            pass
    big = [a for a in amounts if a >= 500]
    if big:
        out["fee.standard"] = _lo(max(big))
        notes.append("price is the largest dollar figure in the document - check it is the "
                     "build price excluding GST, not a total including GST")

    # --- duration
    m = re.search(r"\b(?:(\d{1,2})|(one|two|three|four|five|six|seven|eight|nine|ten|eleven|twelve))"
                  r"[\s-]+weeks?\b", flat, re.I)
    if m:
        word = m.group(2)
        out["timeline.duration"] = _hi(f"{word.capitalize()} weeks" if word else f"{m.group(1)} weeks")

    # --- platform
    for plat in ("WordPress", "Webflow", "Shopify", "Squarespace"):
        if re.search(rf"\b{plat}\b", flat, re.I):
            out["scope.platform"] = _hi(plat)
            break

    # --- contact: a person's name near "attention"/"prepared for"
    m = re.search(r"(?:attention|prepared for|to)\s*:?\s*([A-Z][a-z]+ [A-Z][a-z]+)", flat)
    if m:
        out["client.contact_name"] = _lo(m.group(1))

    # --- pages: lines that look like a page list
    pages = []
    for line in text.split("\n"):
        line = line.strip(" \t|")
        m = re.match(r"^([A-Z][A-Za-z /&'()-]{2,40}?)\s*[|–—:-]\s+(.{6,120})$", line)
        if m and not re.search(r"\$|\d{4}|ABN", line):
            name, purpose = m.group(1).strip(), m.group(2).strip()
            SKIP = {"date", "ref", "subject", "to", "page", "purpose", "atp", "milestone",
                    "description", "amount", "responsibility", "timeline", "milestone / task"}
            if (2 < len(name) < 42
                    and name.lower() not in SKIP
                    and purpose.lower() not in SKIP
                    and "|" not in purpose):
                pages.append([name, purpose])
    if 2 <= len(pages) <= 40:
        out["scope.pages"] = _lo(pages)
        notes.append(f"{len(pages)} possible page rows found - these are guesses, check every one")

    return {"fields": out, "notes": notes, "chars": len(text)}


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("usage: extract.py <proposal.docx|proposal.pdf>")
    print(json.dumps(extract(sys.argv[1]), indent=1))
