import docx, copy, shutil, re, sys

SRC="/Users/asithavajirapani/Downloads/NGA Private - Website ATP and Terms - Encyte.docx"
DST="/Users/asithavajirapani/Downloads/encyte-atp/template/atp-website.docx"
shutil.copy(SRC, DST)
d=docx.Document(DST)

def replace_in_para(p, old, new):
    """Span-aware replace that keeps run formatting where it can."""
    for r in p.runs:                                   # fast path: inside one run
        if old in r.text:
            r.text = r.text.replace(old, new); return True
    full = "".join(r.text for r in p.runs)
    if old not in full: return False
    idx = full.index(old)
    # find run boundaries
    pos, start_run, start_off = 0, None, 0
    for i,r in enumerate(p.runs):
        if pos <= idx < pos+len(r.text):
            start_run, start_off = i, idx-pos; break
        pos += len(r.text)
    end = idx+len(old); pos=0; end_run, end_off = None, 0
    for i,r in enumerate(p.runs):
        if pos < end <= pos+len(r.text):
            end_run, end_off = i, end-pos; break
        pos += len(r.text)
    if start_run is None or end_run is None: return False
    p.runs[start_run].text = p.runs[start_run].text[:start_off] + new
    for i in range(start_run+1, end_run): p.runs[i].text = ""
    if end_run != start_run:
        p.runs[end_run].text = p.runs[end_run].text[end_off:]
    return True

def sub(old, new, limit=None):
    n=0
    targets = list(d.paragraphs) + [c.paragraphs[0] for t in d.tables for r in t.rows for c in r.cells if c.paragraphs]
    for p in targets:
        if limit is not None and n>=limit: break
        if replace_in_para(p, old, new): n+=1
    if n==0: print("  !! NOT FOUND:", old[:70])
    return n

R = [
 # --- parties ---
 ("Encyte Pty Ltd, ABN 62 633 260 474, 175 Maroondah Hwy, Ringwood VIC 3134, Australia",
  "{{provider.legal_name}}, ABN {{provider.abn}}, {{provider.address}}"),
 ("“Encyte”, “we” or “us”", "“{{provider.short_name}}”, “we” or “us”"),
 ("NGA Private Pty Ltd as trustee for the NGA Private Trust, ABN 94 340 745 663, 1/1 International Court, Scoresby VIC 3179",
  "{{client.legal_name}}, ABN {{client.abn}}, {{client.address}}"),
 ("“you” or “NGA Private”", "“you” or “{{client.short_name}}”"),
 ("the new NGA Private website", "the new {{client.short_name}} website"),
 ("proposal 26-NGA-WD-062 dated 8th September 2026", "proposal {{proposal.ref}} dated {{proposal.date}}"),
 ("including proposal 26-NGA-WD-062", "including proposal {{proposal.ref}}"),
 # --- scope ---
 ("A custom-designed website for NGA Private, built on WordPress.",
  "A custom-designed website for {{client.short_name}}, built on {{scope.platform}}."),
 ("copywriting (you supply the page copy), legal and compliance review, SEO campaigns, and a client portal",
  "{{scope.exclusions}}"),
 ("ngaprivate.com.au is registered and paid for by NGA Private.",
  "{{client.domain}} is registered and paid for by {{client.short_name}}."),
 ("a DigitalOcean 2 GB droplet, billed directly to NGA Private by DigitalOcean. Indicatively about USD $14 a month, currently around AUD $20.",
  "a {{hosting.provider}} {{hosting.spec}}, billed directly to {{client.short_name}} by {{hosting.provider}}. {{hosting.cost_note}}"),
 ("(such as WCAG 2.2 Level AA)", "(such as {{scope.accessibility_standard}})"),
 # --- fee ---
 ("The standard build price for this scope is $5,500 excluding GST. NGA Private receives a discount of $650, giving a total of $4,850 excluding GST. Encyte is registered for GST, so GST of $485.00 will be added and shown separately on a valid tax invoice, making the total payable $5,335.00 including GST. Every other amount payable under this agreement – expenses, extra design rounds, changes and interest – is also exclusive of GST.",
  "The standard build price for this scope is {{fee.standard}} excluding {{tax.name}}. {{client.short_name}} receives a discount of {{fee.discount}}, giving a total of {{fee.total}} excluding {{tax.name}}. {{provider.short_name}} is registered for {{tax.name}}, so {{tax.name}} of {{fee.tax}} will be added and shown separately on a valid tax invoice, making the total payable {{fee.total_inc}} including {{tax.name}}. Every other amount payable under this agreement – expenses, extra design rounds, changes and interest – is also exclusive of {{tax.name}}."),
 # --- timeline / rounds / support ---
 ("Eight weeks from kick-off.", "{{timeline.duration}} from kick-off."),
 ("The fee includes two rounds of design revisions before design freeze, and two rounds of feedback during testing and review. Further rounds are charged at $50 per hour plus GST.",
  "The fee includes two rounds of design revisions before design freeze, and two rounds of feedback during testing and review. Further rounds are quoted before they begin and charged only once you approve the quote in writing."),
 ("3.5 Support after launch", "3.5 Technical Support and Maintenance"),
 ("Sixty days of support is included from go-live.", "{{support.included.sentence}} is included from go-live."),
 ("Support hours are 9:00am to 5:00pm Australian Eastern Time, Monday to Friday, excluding public holidays in Victoria.",
  "Support hours are {{support_hours}}."),
 ("After the included 60 days, ongoing care is optional and nothing is required. Our Essential Care plan is $99 plus GST a month and covers security and uptime monitoring, bug fixes, WordPress core and plugin updates, a monthly technical check with a 15-minute call, and email support.",
  "After the included {{support.included.period}}, ongoing technical support and maintenance is optional and nothing is required. Our {{support.plan.name}} plan is {{support.plan.price}} plus {{tax.name}} a {{support.plan.period}} and covers security and uptime monitoring, bug fixes, {{scope.platform}} core and plugin updates, a monthly technical check with a 15-minute call, and email support."),
 ("full WordPress administrator access", "full {{scope.platform}} administrator access"),
 # --- terms / periods ---
 ("payable within 14 days of the invoice date", "payable within {{terms.payment_days}} days of the invoice date"),
 ("at the Reserve Bank of Australia cash rate plus 4% a year", "at the {{interest_benchmark}}"),
 ("cannot go live, for reasons outside our control, for more than 30 days",
  "cannot go live, for reasons outside our control, for more than {{terms.golive_block}} days"),
 ("hear nothing within 5 business days", "hear nothing within {{terms.approval_reminder}} business days"),
 ("still have not heard 3 business days after that", "still have not heard {{terms.approval_escalate}} business days after that"),
 ("for 7 days or more past a due date", "for {{terms.delay_reallocate}} days or more past a due date"),
 ("If we do not hear from you for 30 days", "If we do not hear from you for {{terms.delay_suspend}} days"),
 ("by giving 14 days’ written notice", "by giving {{terms.termination_notice}} days’ written notice"),
 ("within 14 days of a written request", "within {{terms.breach_cure}} days of a written request"),
 ("within 10 business days to try to resolve it", "within {{terms.dispute_meet}} business days to try to resolve it"),
 ("if it is not resolved within 30 days of that notice", "if it is not resolved within {{terms.dispute_mediate}} days of that notice"),
 ("mediation administered by the Resolution Institute under its mediation rules",
  "mediation administered by the {{dispute_body}} under its mediation rules"),
 # --- jurisdiction ---
 ("the Australian Consumer Law and the Privacy Act 1988 (Cth)", "the {{consumer_law}} and the {{privacy_act}}"),
 ("consistent with the Privacy Act 1988 (Cth)", "consistent with the {{privacy_act}}"),
 ("so far as Part IX of the Copyright Act 1968 (Cth) allows", "so far as {{moral_rights_part}} of the {{copyright_act}} allows"),
 ("under the Australian Consumer Law or any other law that cannot be excluded",
  "under the {{consumer_law}} or any other law that cannot be excluded"),
 ("governed by the laws of Victoria, Australia", "governed by the laws of {{governing_law}}"),
 ("NGA Private authorizes Encyte to proceed", "{{client.short_name}} authorizes {{provider.short_name}} to proceed"),
]
print("PARAGRAPH SUBSTITUTIONS")
for old,new in R: print(f"  {sub(old,new)}x  {old[:58]}")

# ---------- tables ----------
def cell(t,r,c,v):
    p=t.rows[r].cells[c].paragraphs[0]
    if p.runs: p.runs[0].text=v; [setattr(x,'text','') for x in p.runs[1:]]
    else: p.add_run(v)

t0=d.tables[0]
cell(t0,1,1,"{{client.contact_name}} – {{client.short_name}}")
cell(t0,1,2,"{{client.contact_name}} – {{client.short_name}}")
cell(t0,2,1,"{{atp.ref}}"); cell(t0,2,2,"{{atp.ref}}")
cell(t0,3,1,"{{atp.subject}}"); cell(t0,3,2,"{{atp.subject}}")
cell(t0,4,0,"Date: {{atp.date}}"); cell(t0,4,1,"Date: {{atp.date}}")
cell(t0,4,2,"{{provider.legal_name}} / {{client.short_name}}")

def keep_one_row(t, first_data, last_data, tokens):
    """Reduce a repeating block to a single template row."""
    for i in range(last_data, first_data, -1):
        t._tbl.remove(t.rows[i]._tr)
    for c,v in enumerate(tokens): cell(t, first_data, c, v)

keep_one_row(d.tables[1], 1, len(d.tables[1].rows)-1, ["{{#pages.0}}", "{{#pages.1}}"])
keep_one_row(d.tables[2], 1, 3, ["{{#milestones.n}}", "{{#milestones.description}}",
                                 "{{#milestones.amount_ex}}", "{{#milestones.amount_inc}}"])
mt=d.tables[2]
last=len(mt.rows)-1
cell(mt,last,2,"{{fee.total}}"); cell(mt,last,3,"{{fee.total_inc}}")
keep_one_row(d.tables[3], 1, len(d.tables[3].rows)-1,
             ["{{#tasks.0}}", "{{#tasks.1}}", "{{#tasks.2}}"])

t4=d.tables[4]
cell(t4,1,3,"{{provider.signatory_name}}")
cell(t4,2,3,"{{provider.signatory_position}}")
cell(t4,3,3,"{{provider.signatory_email}}")
cell(t4,5,3,"")                                    # provider date blank at issue
cell(t4,0,2,"{{provider.legal_name}}"); cell(t4,0,3,"{{provider.legal_name}}")
cell(t4,0,0,"{{client.short_name}}");  cell(t4,0,1,"{{client.short_name}}")

d.save(DST)
print("\nsaved template ->", DST)
