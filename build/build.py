#!/usr/bin/env python3
"""Build an ATP from a jurisdiction pack + an engagement file."""
import docx, json, sys, re, copy, os
from docx.oxml.ns import qn
from decimal import Decimal, ROUND_HALF_UP

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

def money(v):
    q = Decimal(str(v)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return f"${q:,.2f}"

def resolve(v):
    """A string like 'preset:milestones.20-40-40' loads presets/<name>.json."""
    if isinstance(v, str) and v.startswith("preset:"):
        return json.load(open(f"{ROOT}/presets/{v.split(':',1)[1]}.json"))
    return v

# The noun the agreement uses for the thing being delivered. Override per engagement
# with a top-level "deliverable" if a project needs its own word.
DELIVERABLE = {
    "website":    "website",
    "web_app":    "app",
    "mobile_app": "app",
    "brand":      "brand identity",
}

DEFAULTS = {
    "scope.platform": "WordPress",
    "scope.exclusions": "copywriting (you supply the page copy), legal and compliance review, "
                        "SEO campaigns, and a client portal",
    "terms.payment_days": 14,
    "hosting.note": "",
}

def put(o, path, val):
    ks = path.split("."); cur = o
    for k in ks[:-1]: cur = cur.setdefault(k, {})
    cur.setdefault(ks[-1], val)

def load(engagement_file):
    eng = json.load(open(engagement_file))
    jur = json.load(open(f"{ROOT}/jurisdictions/{eng['jurisdiction']}.json"))
    for k, v in DEFAULTS.items(): put(eng, k, v)
    eng["milestones"]          = resolve(eng["milestones"])
    eng["scope"]["pages"]      = resolve(eng["scope"]["pages"])
    eng["scope"]["inclusions"] = resolve(eng["scope"]["inclusions"])
    eng["timeline"]["tasks"]   = resolve(eng["timeline"]["tasks"])
    eng.setdefault("proposal", {}).setdefault("ref", eng["atp"]["ref"])          # defaults to the ATP ref
    eng["atp"].setdefault("subject", f'{eng["client"]["short_name"]} \u2013 Website Design and Build')
    return eng, jur


_WEEK_WORDS = {"one":1,"two":2,"three":3,"four":4,"five":5,"six":6,"seven":7,
               "eight":8,"nine":9,"ten":10,"eleven":11,"twelve":12,"sixteen":16}

def weeks_in(text):
    """Largest week number mentioned in a string, as digits or as a word."""
    best = 0
    for n in re.findall(r"\d+", text or ""):
        best = max(best, int(n))
    for w, n in _WEEK_WORDS.items():
        if re.search(rf"\b{w}\b", (text or "").lower()):
            best = max(best, n)
    return best

# ---------------- validation ----------------
def validate(eng, jur):
    errs = []
    pct = sum(m["percent"] for m in eng["milestones"])
    if pct != 100:
        errs.append(f"milestone percentages sum to {pct}, not 100")
    for m in eng["milestones"]:                       # the Acentura trap
        stated = re.match(r"\s*(\d+)%", m["description"])
        if stated and int(stated.group(1)) != m["percent"]:
            errs.append(f"milestone description says {stated.group(1)}% but percent is {m['percent']}")
    if eng["fee"]["discount"] > eng["fee"]["standard"]:
        errs.append("discount exceeds standard price")
    # the work plan must fit inside the stated duration
    dur = weeks_in(eng.get("timeline", {}).get("duration", ""))
    plan = max([weeks_in(t[-1]) for t in eng.get("timeline", {}).get("tasks", []) if t] or [0])
    if dur and plan and plan > dur:
        errs.append(f"work plan runs to week {plan} but the timeline says {dur} weeks "
                    f"- pick a matching workplan preset, or change the duration")
    if eng.get("engagement_type") != "website":
        errs.append(f"engagement_type '{eng.get('engagement_type')}' has no template yet "
                    f"(website only; web app / mobile app to follow)")
    for path in ["project.name","client.legal_name","client.abn","client.short_name","atp.ref","atp.date",
                 "proposal.ref","scope.platform","timeline.duration"]:
        cur, ok = eng, True
        for k in path.split("."):
            if isinstance(cur, dict) and k in cur and cur[k] not in ("", None): cur = cur[k]
            else: ok = False; break
        if not ok: errs.append(f"missing or empty: {path}")
    if not eng["scope"]["pages"]: errs.append("scope.pages is empty")
    if not eng["scope"].get("inclusions"): errs.append("scope.inclusions is empty")
    if not eng["timeline"].get("tasks"): errs.append("timeline.tasks is empty")
    return errs

# ---------------- derived ----------------
def derive(eng, jur):
    rate = Decimal(str(jur["tax"]["rate"]))
    std  = Decimal(str(eng["fee"]["standard"]))
    disc = Decimal(str(eng["fee"]["discount"]))
    total = std - disc
    tax   = (total * rate).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    d = {
      "deliverable": eng.get("deliverable") or DELIVERABLE.get(eng.get("engagement_type"), "work"),
      "fee.standard": money(std), "fee.discount": money(disc),
      "fee.total": money(total), "fee.tax": money(tax), "fee.total_inc": money(total+tax),
    }
    tx, cl = jur["tax"]["name"], eng["client"]["short_name"]
    d["fee.headline"] = (
        f'The standard build price for this scope is {money(std)} excluding {tx}. '
        f'{cl} receives a discount of {money(disc)}, giving a total of {money(total)} excluding {tx}.'
        if disc > 0 else
        f'The price for this scope is {money(total)} excluding {tx}.')
    rows, run_ex = [], Decimal("0")
    for i,m in enumerate(eng["milestones"]):
        ex = (total * Decimal(str(m["percent"])) / 100).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        if i == len(eng["milestones"])-1: ex = total - run_ex      # last absorbs rounding
        run_ex += ex
        inc = (ex * (1+rate)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        rows.append([str(i+1), m["description"], money(ex), money(inc)])
    assert run_ex == total, f"milestone amounts {run_ex} != total {total}"
    d["_milestone_rows"] = rows
    inc = eng["support"]["included"]
    d["support.included.period"]   = f'{inc["value"]} {inc["unit"]}'
    d["support.included.sentence"] = f'{inc["value"]} {inc["unit"]} of support'
    d["support.plan.price"] = money(eng["support"]["plan"]["price"])
    return d

def flatten(prefix, obj, out):
    for k,v in obj.items():
        key = f"{prefix}{k}"
        if isinstance(v, dict): flatten(key+".", v, out)
        elif isinstance(v, (str,int,float)): out[key] = str(v)
    return out

# ---------------- fill ----------------
def fill_text(doc, tokens):
    def do(p):
        for r in p.runs:
            for k,v in tokens.items():
                t = "{{"+k+"}}"
                if t in r.text: r.text = r.text.replace(t, v)
    def walk(container):
        for p in container.paragraphs: do(p)
        for t in container.tables:
            for row in t.rows:
                for c in row.cells: walk(c)
    walk(doc)
    for s in doc.sections:                      # headers and footers carry tokens too
        for part in (s.header, s.first_page_header, s.even_page_header,
                     s.footer, s.first_page_footer, s.even_page_footer):
            walk(part)

BLOCK_OPEN  = re.compile(r"^\{\{\?(\w+)\}\}$")
BLOCK_CLOSE = re.compile(r"^\{\{/(\w+)\}\}$")

def strip_blocks(doc, engagement_type):
    """{{?website}} ... {{/website}} keeps its contents only for that engagement type.

    Markers are always removed. Everything between them survives only when the name
    matches. Runs before fill_rows, so a stripped table cannot shift table indices.
    """
    children = list(doc.element.body)
    drop, kept, cut, open_at, name = [], 0, 0, None, None
    for idx, el in enumerate(children):
        if el.tag != qn('w:p'):
            continue
        txt = "".join(t.text or "" for t in el.iter(qn('w:t'))).strip()
        m = BLOCK_OPEN.match(txt)
        if m:
            if open_at is not None:
                raise ValueError(f"nested conditional block {{{{?{m.group(1)}}}}} inside {{{{?{name}}}}}")
            open_at, name = idx, m.group(1)
            continue
        m = BLOCK_CLOSE.match(txt)
        if m:
            if open_at is None or m.group(1) != name:
                raise ValueError(f"unmatched {{{{/{m.group(1)}}}}}")
            span = children[open_at:idx + 1]
            if name == engagement_type:
                drop += [span[0], span[-1]]; kept += 1
            else:
                drop += span; cut += 1
            open_at = name = None
    if open_at is not None:
        raise ValueError(f"unclosed conditional block {{{{?{name}}}}}")
    for el in drop:
        el.getparent().remove(el)
    return kept, cut

def find_table(doc, marker):
    """Locate a repeating table by its {{#marker}} rather than by index."""
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                if "{{#" + marker in c.text:
                    return t
    return None

def set_cell(t, r, c, v):
    p = t.rows[r].cells[c].paragraphs[0]
    if p.runs:
        p.runs[0].text = v
        for x in p.runs[1:]: x.text = ""
    else: p.add_run(v)

def fill_rows(table, marker, data):
    """Clone the row carrying {{#marker...}} once per data item."""
    tpl_idx = None
    for i,row in enumerate(table.rows):
        if any("{{#"+marker in c.text for c in row.cells): tpl_idx = i; break
    if tpl_idx is None: return 0
    tpl_tr = copy.deepcopy(table.rows[tpl_idx]._tr)
    for n,item in enumerate(data):
        if n == 0: target = tpl_idx
        else:
            table.rows[tpl_idx+n-1]._tr.addnext(copy.deepcopy(tpl_tr)); target = tpl_idx+n
        for c,val in enumerate(item): set_cell(table, target, c, str(val))
    return len(data)

def build(engagement_file, out_file):
    eng, jur = load(engagement_file)
    errs = validate(eng, jur)
    if errs:
        print("VALIDATION FAILED"); [print("  -", e) for e in errs]; sys.exit(1)
    d = derive(eng, jur)

    tokens = {}
    flatten("", jur, tokens)                 # provider.*, governing_law, tax.name, ...
    flatten("", eng, tokens)                 # client.*, atp.*, scope.platform, terms.*
    tokens.update({k:v for k,v in d.items() if not k.startswith("_")})

    doc = docx.Document(f"{ROOT}/template/atp-website.docx")

    # conditional blocks BEFORE anything addresses a table
    kept, cut = strip_blocks(doc, eng.get("engagement_type"))

    # repeating blocks, located by marker so a stripped table cannot shift indices
    for marker, data in (("pages",      eng["scope"]["pages"]),
                         ("milestones", d["_milestone_rows"]),
                         ("tasks",      eng["timeline"]["tasks"])):
        t = find_table(doc, marker)
        if t is not None:
            fill_rows(t, marker, data)

    # inclusion bullets: clone the marker paragraph
    for p in doc.paragraphs:
        if "{{#inclusion}}" in p.text:
            anchor = p._p                       # advance the anchor, or clones stack in reverse
            for i, text in enumerate(eng["scope"]["inclusions"]):
                if i == 0:
                    tgt = p
                else:
                    clone = copy.deepcopy(p._p)
                    anchor.addnext(clone); anchor = clone
                    tgt = docx.text.paragraph.Paragraph(clone, p._parent)
                for r in tgt.runs: r.text = ""
                tgt.runs[0].text = "\u2022\u2003" + text
            break

    fill_text(doc, tokens)
    doc.save(out_file)

    leftover = set()
    chk = docx.Document(out_file)
    blob = "\n".join(p.text for p in chk.paragraphs) + "\n" + "\n".join(
        c.text for t in chk.tables for r in t.rows for c in r.cells)
    for s_ in chk.sections:                     # check headers/footers as well
        for part in (s_.header, s_.first_page_header, s_.footer, s_.first_page_footer):
            blob += "\n" + "\n".join(p.text for p in part.paragraphs)
    for m in re.findall(r"\{\{[^}]+\}\}", blob): leftover.add(m)
    print(f"built -> {out_file}")
    print(f"  pages {len(eng['scope']['pages'])} · milestones {len(d['_milestone_rows'])} · tasks {len(eng['timeline']['tasks'])}")
    print(f"  type {eng.get('engagement_type')} · deliverable '{d['deliverable']}' · blocks kept {kept}, cut {cut}")
    print(f"  total {d['fee.total']} ex / {d['fee.total_inc']} inc {jur['tax']['name']}")
    print("  UNREPLACED TOKENS:", sorted(leftover) if leftover else "none")
    return not leftover

if __name__ == "__main__":
    ok = build(sys.argv[1], sys.argv[2])
    sys.exit(0 if ok else 2)
